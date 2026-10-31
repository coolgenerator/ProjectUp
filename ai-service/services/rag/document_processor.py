"""
Document Processor
Handles document loading, parsing, and chunking for RAG
"""

import logging
import mimetypes
from typing import List, Dict, Any, Optional
from pathlib import Path
import hashlib

# Document loaders
from langchain_community.document_loaders import (
    PyPDFLoader,
    TextLoader,
    UnstructuredMarkdownLoader,
)
from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter

# Python packages for direct file handling
try:
    from pypdf import PdfReader
    HAS_PYPDF = True
except ImportError:
    HAS_PYPDF = False

try:
    from docx import Document as DocxDocument
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False

logger = logging.getLogger(__name__)


class DocumentProcessor:
    """
    Processes documents for RAG indexing

    Supports:
    - PDF files
    - Word documents (.docx)
    - Markdown (.md)
    - Text files (.txt)
    - Code files (.py, .js, .java, etc.)
    """

    SUPPORTED_EXTENSIONS = {
        ".pdf": "pdf",
        ".txt": "text",
        ".md": "markdown",
        ".markdown": "markdown",
        ".docx": "docx",
        ".py": "code",
        ".js": "code",
        ".ts": "code",
        ".jsx": "code",
        ".tsx": "code",
        ".java": "code",
        ".cpp": "code",
        ".c": "code",
        ".go": "code",
        ".rs": "code",
        ".json": "text",
        ".yaml": "text",
        ".yml": "text",
        ".xml": "text",
        ".html": "text",
        ".css": "text"
    }

    def __init__(
        self,
        chunk_size: int = 1000,
        chunk_overlap: int = 200,
        length_function: Any = len
    ):
        """
        Initialize document processor

        Args:
            chunk_size: Maximum size of each chunk
            chunk_overlap: Overlap between chunks
            length_function: Function to measure text length
        """
        self.chunk_size = chunk_size
        self.chunk_overlap = chunk_overlap

        # Initialize text splitter
        self.text_splitter = RecursiveCharacterTextSplitter(
            chunk_size=chunk_size,
            chunk_overlap=chunk_overlap,
            length_function=length_function,
            separators=["\n\n", "\n", " ", ""]
        )

        logger.info(
            f"Document processor initialized (chunk_size={chunk_size}, "
            f"overlap={chunk_overlap})"
        )

    def is_supported(self, file_path: str) -> bool:
        """
        Check if file type is supported

        Args:
            file_path: Path to file

        Returns:
            Whether file is supported
        """
        ext = Path(file_path).suffix.lower()
        return ext in self.SUPPORTED_EXTENSIONS

    def get_file_type(self, file_path: str) -> Optional[str]:
        """
        Get file type category

        Args:
            file_path: Path to file

        Returns:
            File type category or None
        """
        ext = Path(file_path).suffix.lower()
        return self.SUPPORTED_EXTENSIONS.get(ext)

    def load_document(
        self,
        file_path: str,
        metadata: Optional[Dict[str, Any]] = None
    ) -> List[Document]:
        """
        Load and parse a document

        Args:
            file_path: Path to document
            metadata: Optional metadata to attach

        Returns:
            List of Document objects
        """
        file_path = Path(file_path)

        if not file_path.exists():
            raise FileNotFoundError(f"File not found: {file_path}")

        if not self.is_supported(str(file_path)):
            raise ValueError(
                f"Unsupported file type: {file_path.suffix}. "
                f"Supported: {list(self.SUPPORTED_EXTENSIONS.keys())}"
            )

        file_type = self.get_file_type(str(file_path))
        logger.info(f"Loading {file_type} document: {file_path.name}")

        # Base metadata
        base_metadata = {
            "source": str(file_path),
            "file_name": file_path.name,
            "file_type": file_type,
            "file_size": file_path.stat().st_size,
            "file_hash": self._compute_file_hash(file_path)
        }

        if metadata:
            base_metadata.update(metadata)

        # Load based on file type
        try:
            if file_type == "pdf":
                documents = self._load_pdf(file_path, base_metadata)
            elif file_type == "docx":
                documents = self._load_docx(file_path, base_metadata)
            elif file_type == "markdown":
                documents = self._load_markdown(file_path, base_metadata)
            elif file_type in ["text", "code"]:
                documents = self._load_text(file_path, base_metadata)
            else:
                raise ValueError(f"Unsupported file type: {file_type}")

            logger.info(f"Loaded {len(documents)} pages/sections from {file_path.name}")
            return documents

        except Exception as e:
            logger.error(f"Error loading document {file_path}: {e}")
            raise

    def _load_pdf(self, file_path: Path, metadata: Dict[str, Any]) -> List[Document]:
        """Load PDF document"""
        if not HAS_PYPDF:
            raise ImportError("pypdf is required for PDF processing")

        try:
            loader = PyPDFLoader(str(file_path))
            documents = loader.load()

            # Add metadata to each page
            for i, doc in enumerate(documents):
                doc.metadata.update(metadata)
                doc.metadata["page"] = i + 1

            return documents
        except Exception as e:
            logger.error(f"Error loading PDF: {e}")
            # Fallback: try direct pypdf
            return self._load_pdf_direct(file_path, metadata)

    def _load_pdf_direct(self, file_path: Path, metadata: Dict[str, Any]) -> List[Document]:
        """Load PDF using pypdf directly"""
        documents = []
        reader = PdfReader(str(file_path))

        for i, page in enumerate(reader.pages):
            text = page.extract_text()
            if text.strip():
                doc_metadata = {**metadata, "page": i + 1}
                documents.append(Document(page_content=text, metadata=doc_metadata))

        return documents

    def _load_docx(self, file_path: Path, metadata: Dict[str, Any]) -> List[Document]:
        """Load Word document"""
        if not HAS_DOCX:
            raise ImportError("python-docx is required for DOCX processing")

        doc = DocxDocument(str(file_path))
        text_content = []

        for para in doc.paragraphs:
            if para.text.strip():
                text_content.append(para.text)

        full_text = "\n\n".join(text_content)

        return [Document(page_content=full_text, metadata=metadata)]

    def _load_markdown(self, file_path: Path, metadata: Dict[str, Any]) -> List[Document]:
        """Load Markdown document"""
        try:
            loader = UnstructuredMarkdownLoader(str(file_path))
            documents = loader.load()

            for doc in documents:
                doc.metadata.update(metadata)

            return documents
        except Exception as e:
            logger.warning(f"UnstructuredMarkdownLoader failed, using TextLoader: {e}")
            return self._load_text(file_path, metadata)

    def _load_text(self, file_path: Path, metadata: Dict[str, Any]) -> List[Document]:
        """Load text/code file"""
        with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
            text = f.read()

        return [Document(page_content=text, metadata=metadata)]

    def chunk_documents(
        self,
        documents: List[Document],
        add_chunk_metadata: bool = True
    ) -> List[Document]:
        """
        Split documents into chunks

        Args:
            documents: List of documents to chunk
            add_chunk_metadata: Whether to add chunk-specific metadata

        Returns:
            List of chunked documents
        """
        all_chunks = []

        for doc_idx, doc in enumerate(documents):
            chunks = self.text_splitter.split_documents([doc])

            if add_chunk_metadata:
                for chunk_idx, chunk in enumerate(chunks):
                    chunk.metadata["chunk_index"] = chunk_idx
                    chunk.metadata["total_chunks"] = len(chunks)
                    chunk.metadata["chunk_id"] = f"{doc_idx}_{chunk_idx}"

            all_chunks.extend(chunks)

        logger.info(f"Split {len(documents)} documents into {len(all_chunks)} chunks")
        return all_chunks

    def process_document(
        self,
        file_path: str,
        metadata: Optional[Dict[str, Any]] = None,
        chunk: bool = True
    ) -> List[Document]:
        """
        Complete document processing pipeline

        Args:
            file_path: Path to document
            metadata: Optional metadata
            chunk: Whether to chunk the document

        Returns:
            List of processed document chunks
        """
        # Load document
        documents = self.load_document(file_path, metadata)

        # Chunk if requested
        if chunk:
            documents = self.chunk_documents(documents)

        return documents

    def process_multiple_documents(
        self,
        file_paths: List[str],
        metadata: Optional[Dict[str, Any]] = None,
        chunk: bool = True
    ) -> List[Document]:
        """
        Process multiple documents

        Args:
            file_paths: List of file paths
            metadata: Optional metadata for all documents
            chunk: Whether to chunk documents

        Returns:
            List of all processed chunks
        """
        all_documents = []

        for file_path in file_paths:
            try:
                # Add file-specific metadata
                file_metadata = metadata.copy() if metadata else {}
                file_metadata["document_id"] = self._compute_file_hash(Path(file_path))

                docs = self.process_document(file_path, file_metadata, chunk)
                all_documents.extend(docs)

            except Exception as e:
                logger.error(f"Error processing {file_path}: {e}")
                continue

        logger.info(f"Processed {len(file_paths)} documents into {len(all_documents)} chunks")
        return all_documents

    def extract_text_from_content(
        self,
        content: str,
        metadata: Optional[Dict[str, Any]] = None,
        chunk: bool = True
    ) -> List[Document]:
        """
        Process text content directly (e.g., from API upload)

        Args:
            content: Text content
            metadata: Optional metadata
            chunk: Whether to chunk the content

        Returns:
            List of document chunks
        """
        doc = Document(page_content=content, metadata=metadata or {})

        if chunk:
            return self.text_splitter.split_documents([doc])
        else:
            return [doc]

    def _compute_file_hash(self, file_path: Path) -> str:
        """Compute SHA-256 hash of file"""
        sha256_hash = hashlib.sha256()
        with open(file_path, "rb") as f:
            for byte_block in iter(lambda: f.read(4096), b""):
                sha256_hash.update(byte_block)
        return sha256_hash.hexdigest()[:16]  # First 16 chars

    def get_document_summary(self, documents: List[Document]) -> Dict[str, Any]:
        """Get summary statistics about processed documents"""
        total_chars = sum(len(doc.page_content) for doc in documents)
        sources = set(doc.metadata.get("source", "unknown") for doc in documents)

        return {
            "total_chunks": len(documents),
            "total_characters": total_chars,
            "unique_sources": len(sources),
            "sources": list(sources),
            "average_chunk_size": total_chars // len(documents) if documents else 0
        }
